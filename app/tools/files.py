from pathlib import Path
from ..config import SEARCH_ROOTS, MAX_FILE_BYTES, MAX_TEXT_CHARS

def find_files(query:str,limit:int=50):
    q=query.lower().strip(); results=[]
    for root in SEARCH_ROOTS:
        p=Path(root)
        if not p.exists(): continue
        try:
            for item in p.rglob('*'):
                if len(results)>=limit: return results
                if item.is_file() and q in item.name.lower(): results.append(str(item))
        except (PermissionError,OSError): continue
    return results

def read_file(path:str):
    p=Path(path)
    if not p.is_file(): raise FileNotFoundError(str(p))
    if p.stat().st_size>MAX_FILE_BYTES: raise ValueError('File quá lớn; hãy chỉ định phần cần đọc.')
    ext=p.suffix.lower()
    if ext in {'.txt','.md','.csv','.json','.log','.py','.js','.html','.css'}:
        return p.read_text(encoding='utf-8',errors='ignore')[:MAX_TEXT_CHARS]
    if ext=='.pdf':
        from pypdf import PdfReader
        reader=PdfReader(str(p)); text='\n'.join(page.extract_text() or '' for page in reader.pages)
        return text[:MAX_TEXT_CHARS]
    if ext=='.docx':
        from docx import Document
        doc=Document(str(p)); return '\n'.join(x.text for x in doc.paragraphs)[:MAX_TEXT_CHARS]
    if ext=='.xlsx':
        from openpyxl import load_workbook
        wb=load_workbook(str(p),read_only=True,data_only=True); chunks=[]
        for ws in wb.worksheets:
            chunks.append(f'## Sheet: {ws.title}')
            for row in ws.iter_rows(values_only=True): chunks.append(' | '.join('' if v is None else str(v) for v in row))
            if sum(len(x) for x in chunks)>MAX_TEXT_CHARS: break
        return '\n'.join(chunks)[:MAX_TEXT_CHARS]
    raise ValueError(f'Chưa hỗ trợ định dạng {ext}')
